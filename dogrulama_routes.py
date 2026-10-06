"""
DYS — Doğrulama (D04 F03 Mastar & Cihaz Doğrulama Kaydı, 2026): mastar / cihaz kartı + periyodik doğrulama oturumları (referans → ölçülen → sapma / karar otomatik).
  /dogrulama                  takip listesi: son / gelecek doğrulama, kalan gün, karar; yeni kart; Excel
  /dogrulama/<id>             kart (üst bilgiler, görseller), doğrulama geçmişi, yeni doğrulama girişi (yalnız ölçülen değer)
  /dogrulama/<id>/yazdir      D04 F03 çıktısı (A4)
  /dogrulama/gorsel/<gid>     görsel
Kalibrasyon listesiyle bağı şimdilik yoktur (ekipman_id alanı ileride). Yetki: kart / görsel / silme Admin / Doküman Kontrol; doğrulama girişi Sadece Görüntüleme hariç herkes.
Karar: |ölçülen − referans| ≤ |tolerans| (satır toleransı, yoksa kartın izin verilen sapması) → OK, değilse NOK; tolerans tanımsızsa karar boş.
"""
import calendar
import io
import os
from datetime import date, datetime

from flask import render_template, request, redirect, url_for, flash, abort, session, send_file
from werkzeug.utils import secure_filename

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from config import Config  # noqa: E402
from models import DogrulamaKarti, DogrulamaOturumu, DogrulamaOlcum, DogrulamaGorsel, User  # noqa: E402
from helpers import resolve_file_path  # noqa: E402

TURLER = ("Mastar", "Cihaz", "Fikstür", "Terazi", "Termokupl")
RESIM = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def _yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _yazabilir():
    return session.get("rol") != "Sadece Görüntüleme"


def _k(s):
    return " ".join((s or "").replace("İ", "i").replace("I", "ı").lower().split())


def _sayi(v):
    v = (str(v) if v is not None else "").strip().replace(",", ".")
    if not v:
        return None
    try:
        return float(v)
    except ValueError:
        raise ValueError(v)


def _ay_ekle(t, n):
    if not t or not n:
        return None
    y, m = divmod(t.month - 1 + n, 12)
    return date(t.year + y, m + 1, min(t.day, calendar.monthrange(t.year + y, m + 1)[1]))


def _karar(sonuc, ref, tol, kart_tol):
    if sonuc is None or ref is None:
        return None
    t = tol if tol is not None else kart_tol
    if t is None:
        return None
    return "OK" if abs(round(sonuc - ref, 6)) <= abs(t) + 1e-9 else "NOK"


def _ozet(db, kartlar):
    """kart_id → {son, gelecek, kalan, karar, adet, ileri}: son = bugüne kadarki en son oturum."""
    bugun = datetime.now()
    out = {}
    for o in db.query(DogrulamaOturumu).order_by(DogrulamaOturumu.tarih).all():
        d = out.setdefault(o.kart_id, {"son": None, "karar": None, "adet": 0, "ileri": 0, "yapan": None})
        d["adet"] += 1
        if o.tarih and o.tarih > bugun:
            d["ileri"] += 1
        elif o.tarih:
            d["son"], d["karar"], d["yapan"] = o.tarih, o.karar, o.yapan
    for k in kartlar:
        d = out.setdefault(k.id, {"son": None, "karar": None, "adet": 0, "ileri": 0, "yapan": None})
        d["gelecek"] = _ay_ekle(d["son"].date(), k.dogrulama_ay) if d["son"] and k.dogrulama_ay else None
        d["kalan"] = (d["gelecek"] - date.today()).days if d["gelecek"] else None
    return out


def _siniflar(d):
    """Öncelik sınıfı: 0 geciken (≤ 1 yıl) · 5 1 yıldan eski gecikme (uzun süredir doğrulanmıyor) · 1 bugün / 7 gün · 2 30 gün içinde · 3 normal · 4 doğrulama kaydı yok."""
    k = d.get("kalan")
    if k is None:
        return 4
    return 5 if k < -365 else 0 if k < 0 else 1 if k <= 7 else 2 if k <= 30 else 3


def _etiket(d):
    k = d.get("kalan")
    if k is None:
        return "Doğrulanmadı" if not d.get("son") else "Sıklık tanımsız"
    return f"{-k} gün gecikti" if k < 0 else ("Bugün" if k == 0 else f"{k} gün kaldı")


def _aktif_liste(db):
    kartlar = db.query(DogrulamaKarti).filter(DogrulamaKarti.aktif.isnot(False)).order_by(DogrulamaKarti.grup, DogrulamaKarti.no).all()
    oz = _ozet(db, kartlar)
    return kartlar, oz


# ─────────────────────────── liste ───────────────────────────
@app.route("/dogrulama", methods=["GET", "POST"])
@login_required
def dogrulama_liste():
    db = get_db()
    try:
        if request.method == "POST":
            if not _yonetici():
                abort(403)
            f = request.form
            no, ad = (f.get("no") or "").strip(), (f.get("ad") or "").strip()
            if not (no and ad):
                flash("No ve ad zorunludur.", "error")
                return redirect(url_for("dogrulama_liste"))
            try:
                ms = _sayi(f.get("max_sapma"))
            except ValueError:
                flash("İzin verilen sapma sayı olmalıdır.", "error")
                return redirect(url_for("dogrulama_liste"))
            k = DogrulamaKarti(no=no[:80], ad=ad[:250], tur=f.get("tur") if f.get("tur") in TURLER else "Mastar", grup=(f.get("grup") or "").strip()[:80] or None,
                               dogrulama_ay=f.get("dogrulama_ay", type=int), max_sapma=abs(ms) if ms is not None else None, aktif=True, durum="Aktif",
                               kullanim_yeri=(f.get("kullanim_yeri") or "").strip()[:150] or None, kaynak="elle")
            db.add(k); db.commit()
            log_action(db, "Oluşturma", detay=f"Doğrulama kartı: {k.no} — {k.ad}")
            return redirect(url_for("dogrulama_kart", kid=k.id))
        f = request.args
        q, grup, tur, durum, filtre = _k(f.get("q")), f.get("grup", ""), f.get("tur", ""), f.get("durum", "aktif"), f.get("filtre", "")
        kartlar = db.query(DogrulamaKarti).order_by(DogrulamaKarti.grup, DogrulamaKarti.no).all()
        oz = _ozet(db, kartlar)
        gruplar = sorted({k.grup for k in kartlar if k.grup}, key=_k)
        taban, grup_say = [], {}
        for k in kartlar:
            if durum == "aktif" and k.aktif is False or durum == "pasif" and k.aktif is not False:
                continue
            if tur and k.tur != tur or q and q not in _k(f"{k.no} {k.ad} {k.urun} {k.grup} {k.kullanim_yeri}"):
                continue
            taban.append((k, oz[k.id]))
            grup_say[k.grup or "—"] = grup_say.get(k.grup or "—", 0) + 1
        L = [(k, d) for k, d in taban if not grup or k.grup == grup]
        sayac = {"toplam": len(L), "geciken": sum(1 for k, d in L if _siniflar(d) == 0), "eski": sum(1 for k, d in L if _siniflar(d) == 5), "hafta": sum(1 for k, d in L if _siniflar(d) == 1),
                 "ay": sum(1 for k, d in L if _siniflar(d) in (1, 2)), "yok": sum(1 for k, d in L if _siniflar(d) == 4), "nok": sum(1 for k, d in L if d["karar"] == "NOK")}
        if filtre == "geciken":
            L = [(k, d) for k, d in L if _siniflar(d) == 0]
        elif filtre == "eski":
            L = [(k, d) for k, d in L if _siniflar(d) == 5]
        elif filtre == "hafta":
            L = [(k, d) for k, d in L if _siniflar(d) == 1]
        elif filtre == "ay":
            L = [(k, d) for k, d in L if _siniflar(d) in (1, 2)]
        elif filtre == "yok":
            L = [(k, d) for k, d in L if _siniflar(d) == 4]
        elif filtre == "nok":
            L = [(k, d) for k, d in L if d["karar"] == "NOK"]
        # tarihi geçenler en başta (en çok gecikenden), sonra yaklaşanlar; doğrulama kaydı olmayanlar sonda
        sira = {0: 0, 5: 1, 1: 2, 2: 3, 3: 4, 4: 5}
        L.sort(key=lambda x: (sira[_siniflar(x[1])], x[1]["kalan"] if x[1]["kalan"] is not None else 10 ** 6, _k(x[0].no)))
        if f.get("excel") == "1":
            return _excel(L)
        return render_template("dogrulama_liste.html", liste=L, gruplar=gruplar, grup_say=grup_say, q=f.get("q", ""), grup=grup, tur=tur, durum=durum, filtre=filtre,
                               sayac=sayac, yonetici=_yonetici(), turler=TURLER, etiket=_etiket, sinif=_siniflar)
    finally:
        db.close()


def _excel(L):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Doğrulama Takip"
    ws.append(["No", "Ad", "Tür", "Grup", "Ürün / takım", "Sıklık (ay)", "Son doğrulama", "Gelecek doğrulama", "Kalan gün", "Son karar", "Durum", "Kullanım yeri", "Önem"])
    for k, d in L:
        ws.append([k.no, k.ad, k.tur, k.grup, k.urun, k.dogrulama_ay, d["son"].strftime("%d.%m.%Y") if d["son"] else None, d["gelecek"].strftime("%d.%m.%Y") if d["gelecek"] else None,
                   d["kalan"], d["karar"], "Aktif" if k.aktif is not False else "Pasif", k.kullanim_yeri, _etiket(d)])
    for c, w in zip("ABCDEFGHIJKLM", (22, 46, 9, 14, 26, 10, 14, 16, 10, 10, 9, 22, 16)):
        ws.column_dimensions[c].width = w
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return send_file(b, as_attachment=True, download_name="dogrulama_takip.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ─────────────────────────── rapor ───────────────────────────
@app.route("/dogrulama/rapor")
@login_required
def dogrulama_rapor():
    """Önce tarihi geçenler (en çok gecikenden), sonra 30 gün içinde dolacaklar, son kararı NOK olanlar, hiç doğrulanmamışlar."""
    db = get_db()
    try:
        kartlar, oz = _aktif_liste(db)
        grup = request.args.get("grup", "")
        P = [(k, oz[k.id]) for k in kartlar if not grup or k.grup == grup]
        gec = sorted([x for x in P if _siniflar(x[1]) == 0], key=lambda x: x[1]["kalan"])
        eski = sorted([x for x in P if _siniflar(x[1]) == 5], key=lambda x: x[1]["kalan"], reverse=True)
        yak = sorted([x for x in P if _siniflar(x[1]) in (1, 2)], key=lambda x: x[1]["kalan"])
        nok = [x for x in P if x[1]["karar"] == "NOK"]
        yok = [x for x in P if _siniflar(x[1]) == 4]
        ozet = {}
        for k, d in P:
            g = ozet.setdefault(k.grup or "—", {"toplam": 0, "geciken": 0, "eski": 0, "ay": 0, "nok": 0, "yok": 0})
            g["toplam"] += 1
            g["geciken"] += _siniflar(d) == 0
            g["eski"] += _siniflar(d) == 5
            g["ay"] += _siniflar(d) in (1, 2)
            g["nok"] += d["karar"] == "NOK"
            g["yok"] += _siniflar(d) == 4
        if request.args.get("excel") == "1":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook(); wb.remove(wb.active)
            for ad, L, renk in (("Gecikenler", gec, "FDE2E2"), ("1 yıldan eski gecikme", eski, "E5E7EB"), ("30 gün içinde", yak, "FFF1C2"), ("Son karar NOK", nok, "FDE2E2"), ("Doğrulanmamış", yok, "EEEEEE")):
                ws = wb.create_sheet(ad)
                ws.append(["No", "Ad", "Tür", "Grup", "Ürün / takım", "Sıklık (ay)", "Son doğrulama", "Gelecek doğrulama", "Gecikme / kalan gün", "Son karar", "Son yapan", "Kullanım yeri"])
                for c in ws[1]:
                    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor=renk)
                for k, d in L:
                    ws.append([k.no, k.ad, k.tur, k.grup, k.urun, k.dogrulama_ay, d["son"].strftime("%d.%m.%Y") if d["son"] else None, d["gelecek"].strftime("%d.%m.%Y") if d["gelecek"] else None,
                               (-d["kalan"] if d["kalan"] is not None and d["kalan"] < 0 else d["kalan"]), d["karar"], d["yapan"], k.kullanim_yeri])
                for col, w in zip("ABCDEFGHIJKL", (22, 46, 9, 14, 26, 10, 14, 16, 16, 10, 16, 22)):
                    ws.column_dimensions[col].width = w
            b = io.BytesIO(); wb.save(b); b.seek(0)
            return send_file(b, as_attachment=True, download_name=f"dogrulama_rapor_{date.today():%Y%m%d}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        gruplar = sorted({k.grup for k in kartlar if k.grup}, key=_k)
        return render_template("dogrulama_rapor.html", gec=gec, eski=eski, yak=yak, nok=nok, yok=yok, ozet=ozet, gruplar=gruplar, grup=grup, bugun=date.today(), toplam=len(P))
    finally:
        db.close()


# ─────────────────────────── kart ───────────────────────────
def _kart_verisi(db, k):
    ot = db.query(DogrulamaOturumu).filter_by(kart_id=k.id).order_by(DogrulamaOturumu.tarih.desc(), DogrulamaOturumu.sira.desc()).all()
    olc = {}
    for m in db.query(DogrulamaOlcum).filter(DogrulamaOlcum.oturum_id.in_([o.id for o in ot] or [0])).order_by(DogrulamaOlcum.oturum_id, DogrulamaOlcum.sira).all():
        olc.setdefault(m.oturum_id, []).append(m)
    gor = db.query(DogrulamaGorsel).filter_by(kart_id=k.id).order_by(DogrulamaGorsel.tur, DogrulamaGorsel.id).all()
    return ot, olc, gor, _ozet(db, [k])[k.id]


@app.route("/dogrulama/<int:kid>", methods=["GET", "POST"])
@login_required
def dogrulama_kart(kid):
    db = get_db()
    try:
        k = db.get(DogrulamaKarti, kid)
        if not k:
            abort(404)
        git = url_for("dogrulama_kart", kid=k.id)
        if request.method == "POST":
            f = request.form
            a = f.get("action")
            if a == "oturum":
                if not _yazabilir():
                    abort(403)
                return _oturum_ekle(db, k, f, git)
            if not _yonetici():
                abort(403)
            if a == "bilgi":
                for alan, n in (("no", 80), ("ad", 250), ("grup", 80), ("urun", 120), ("devreye_alma_olcusu", 200), ("kalibrasyon_siklik", 40), ("olcum_araligi", 100), ("marka_model", 200),
                                ("ortam", 100), ("mastar_takimi", 300), ("kullanim_yeri", 150), ("yontem", 1000), ("durum", 60)):
                    v = (f.get(alan) or "").strip()[:n]
                    if alan in ("no", "ad") and not v:
                        continue
                    setattr(k, alan, v or None)
                k.tur = f.get("tur") if f.get("tur") in TURLER else k.tur
                k.dogrulama_ay = f.get("dogrulama_ay", type=int)
                try:
                    ms = _sayi(f.get("max_sapma"))
                    k.max_sapma = abs(ms) if ms is not None else None
                except ValueError:
                    flash("İzin verilen sapma sayı olmalıdır.", "error")
                try:
                    k.devreye_alma_tarihi = datetime.strptime(f.get("devreye_alma_tarihi"), "%Y-%m-%d").date() if f.get("devreye_alma_tarihi") else None
                except ValueError:
                    pass
                k.aktif = f.get("aktif") == "1"
                db.commit(); flash("Kart güncellendi.", "success")
                log_action(db, "Güncelleme", detay=f"Doğrulama kartı güncellendi: {k.no}")
            elif a == "oturum_sil":
                o = db.get(DogrulamaOturumu, f.get("oid", type=int) or 0)
                if o and o.kart_id == k.id:
                    db.query(DogrulamaOlcum).filter_by(oturum_id=o.id).delete(); db.delete(o); db.commit()
                    log_action(db, "Silme", detay=f"Doğrulama oturumu silindi: {k.no}"); flash("Oturum silindi.", "success")
            elif a == "gorsel":
                dosya = request.files.get("dosya")
                if not dosya or not dosya.filename or os.path.splitext(dosya.filename)[1].lower() not in RESIM:
                    flash("PNG / JPG / GIF / WEBP görsel seçin.", "error")
                else:
                    ad = secure_filename(dosya.filename) or "gorsel.png"
                    kd = os.path.join(Config.UPLOAD_FOLDER, "dogrulama", str(k.id)); os.makedirs(kd, exist_ok=True)
                    tur = f.get("tur") if f.get("tur") in ("cihaz", "yontem") else "yontem"
                    ad = f"{tur}_{datetime.now():%Y%m%d%H%M%S}_{ad}"[:190]
                    dosya.save(os.path.join(kd, ad))
                    db.add(DogrulamaGorsel(kart_id=k.id, tur=tur, dosya_adi=ad, dosya_yolu=os.path.join("dogrulama", str(k.id), ad))); db.commit()
                    flash("Görsel eklendi.", "success")
            elif a == "gorsel_sil":
                g = db.get(DogrulamaGorsel, f.get("gid", type=int) or 0)
                if g and g.kart_id == k.id:
                    db.delete(g); db.commit(); flash("Görsel silindi.", "success")
            return redirect(git)
        ot, olc, gor, oz = _kart_verisi(db, k)
        son = ot[0] if ot else None
        hazir = [(m.etiket, m.referans, m.tolerans) for m in (olc.get(next((o.id for o in ot if o.tarih and o.tarih <= datetime.now()), None), []) or [])]
        kullanici = db.get(User, session.get("user_id"))
        return render_template("dogrulama_kart.html", k=k, oturumlar=ot, olc=olc, gorseller=gor, oz=oz, hazir=hazir, yonetici=_yonetici(), yazabilir=_yazabilir(), turler=TURLER,
                               kullanici_adi=(kullanici.ad_soyad if kullanici else ""), simdi=datetime.now().strftime("%Y-%m-%dT%H:%M"), son=son)
    finally:
        db.close()


def _oturum_ekle(db, k, f, git):
    try:
        t = datetime.strptime(f.get("tarih"), "%Y-%m-%dT%H:%M")
    except (ValueError, TypeError):
        flash("Tarih geçersiz.", "error")
        return redirect(git)
    et, rf, tl, sn = f.getlist("etiket"), f.getlist("referans"), f.getlist("tolerans"), f.getlist("sonuc")
    satirlar = []
    try:
        for i in range(len(sn)):
            if not (sn[i] or "").strip():
                continue
            r, tol, s = _sayi(rf[i]) if i < len(rf) else None, _sayi(tl[i]) if i < len(tl) else None, _sayi(sn[i])
            satirlar.append(((et[i] if i < len(et) else "").strip()[:150] or None, r, tol, s))
    except ValueError as e:
        flash(f"Sayı olmayan değer: {e}", "error")
        return redirect(git)
    if not satirlar:
        flash("En az bir ölçülen değer girin.", "error")
        return redirect(git)
    sira = (db.query(DogrulamaOturumu).filter_by(kart_id=k.id).count()) + 1
    o = DogrulamaOturumu(kart_id=k.id, sira=sira, tarih=t, yapan=(f.get("yapan") or "").strip()[:100] or None, yapan_id=session.get("user_id"), ortam=(f.get("ortam") or k.ortam or "")[:100] or None,
                         cihaz=(f.get("cihaz") or "").strip()[:150] or None, notu=(f.get("notu") or "").strip()[:500] or None)
    db.add(o); db.flush()
    kar = []
    for i, (e, r, tol, s) in enumerate(satirlar, 1):
        kr = _karar(s, r, tol, k.max_sapma)
        kar.append(kr)
        db.add(DogrulamaOlcum(oturum_id=o.id, sira=i, etiket=e, referans=r, tolerans=tol, sonuc=s, sapma=round(s - r, 6) if r is not None else None, karar=kr))
    o.karar = "NOK" if "NOK" in kar else ("OK" if "OK" in kar else None)
    db.commit()
    log_action(db, "Oluşturma", detay=f"Doğrulama kaydı: {k.no} — {o.karar or 'karar yok'}")
    if o.karar == "NOK":
        flash("Doğrulama NOK: cihaz / mastar kullanımdan kaldırılmalı, kalibrasyon / düzeltme gerekir; ilgili ölçümler için Uygun Olmayan Ürün değerlendirmesi yapın.", "warning")
    else:
        flash("Doğrulama kaydedildi.", "success")
    return redirect(git)


@app.route("/dogrulama/<int:kid>/yazdir")
@login_required
def dogrulama_yazdir(kid):
    db = get_db()
    try:
        k = db.get(DogrulamaKarti, kid)
        if not k:
            abort(404)
        ot, olc, gor, oz = _kart_verisi(db, k)
        return render_template("dogrulama_yazdir.html", k=k, oturumlar=sorted(ot, key=lambda o: (o.tarih or datetime.min, o.sira)), olc=olc, gorseller=gor, oz=oz)
    finally:
        db.close()


@app.route("/dogrulama/gorsel/<int:gid>")
@login_required
def dogrulama_gorsel(gid):
    db = get_db()
    try:
        g = db.get(DogrulamaGorsel, gid)
        if not g:
            abort(404)
        yol = resolve_file_path(g.dosya_yolu, g.dosya_adi)
        if not yol or not os.path.isfile(yol):
            abort(404)
        return send_file(yol)
    finally:
        db.close()
