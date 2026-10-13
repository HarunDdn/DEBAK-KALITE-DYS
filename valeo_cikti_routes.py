"""
DYS — Valeo formatında ANTETLİ ÇIKTILAR (2026): ürün prosesi (PPAP) verisinden Flow Chart, PFMEA ve Control Plan A4 çıktısı (yazdır / PDF) ve Excel.
  /urun-proses/<id>/cikti/akis             Production Flow Chart (A4 dikey)
  /urun-proses/<id>/cikti/fmea             Process FMEA — operasyon başına sayfa (A4 yatay)
  /urun-proses/<id>/cikti/kontrol-plani    Control Plan (A4 yatay) + Reaction Plan Definitions
  …?excel=1                                Excel çıktısı (Debak logolu)
Her sayfada Debak logolu antet. Veri: ProsesFmeaSatiri / ProsesKpSatiri sistem alanları; Valeo ham verisi (valeo_json) varsa yazdırmada ham puan / alanlar kullanılır
(satır sistemde düzenlenmişse güncel değerler gösterilir). AP yalnız GSI matrisinden hesaplanır (fmea_core); RPN bilgidir.
"""
from sqlalchemy import case as _sa_case
import io
import json
import os
import re
from datetime import date

from flask import render_template, request, abort, send_file, session

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import UrunProsesi, UrunProsesAdimi, ProsesAdimTanim, ProsesFmeaSatiri, ProsesKpSatiri  # noqa: E402
import fmea_core as FC  # noqa: E402
import urun_proses_routes as UP  # noqa: E402

LOGO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "branding", "debak-logo.png")


def _j(s):
    try:
        return json.loads(s) if s else {}
    except ValueError:
        return {}


def _puan(r, vj, anahtar):
    """Yazdırılacak puan: ham değer, sistemdeki (normalize) değer ile uyumluysa ham; aksi halde sistem değeri."""
    db = getattr(r, anahtar, None)
    ham = vj.get(anahtar)
    if ham is not None and FC.normalize({anahtar: ham}, "P").get(anahtar) == db:
        return ham
    return db


def _veri(db, u):
    """Çıktı için ürün verisi: ops[] = {ad, vj, fmea[], kp[]}"""
    katalog = {k.id: k.ad for k in db.query(ProsesAdimTanim).all()}
    adimlar = db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).order_by(UrunProsesAdimi.sira, UrunProsesAdimi.id).all()
    ops = []
    for a in adimlar:
        vj = _j(a.valeo_json)
        std = UP._std(u)
        q = db.query(ProsesFmeaSatiri).filter(ProsesFmeaSatiri.adim_tanim_id == a.adim_tanim_id, (ProsesFmeaSatiri.urun_proses_id == u.id) | (ProsesFmeaSatiri.urun_proses_id.is_(None)))
        q = UP._std_sorgu(q, std).order_by(_sa_case((ProsesFmeaSatiri.urun_proses_id.isnot(None), 1), else_=0), ProsesFmeaSatiri.sira, ProsesFmeaSatiri.id)
        fm = []
        for r in q.all():
            v = _j(r.valeo_json)
            S, O, D = _puan(r, v, "S"), _puan(r, v, "O"), _puan(r, v, "D")
            h = FC.hesapla({"S": r.S, "O": r.O, "D": r.D, "O2": r.O2, "D2": r.D2, "act": r.act}, "P", std)
            rpn = (S * O * D) if (S and O and D) else v.get("rpn")
            O2, D2 = _puan(r, v, "O2"), _puan(r, v, "D2")
            S2 = v.get("S2") if (v.get("S2") is not None and (O2 or D2)) else (S if (O2 or D2) else None)
            rpn2 = (S2 * (O2 or O or 0) * (D2 or D or 0)) if (S2 and (O2 or D2)) else v.get("rpn2")
            fm.append({"fn": r.fn or "", "req": v.get("req", ""), "mode": r.mode or "", "eff": r.eff or "", "S": S, "sinif": v.get("sinif", ""), "cause": r.cause or "", "prev": r.prev or "", "O": O,
                       "det": r.det or "", "D": D, "rpn": rpn, "ap": h["ap"], "act": (r.act if (r.act or not v) else ""), "sorumlu": v.get("sorumlu", ""), "alinan": v.get("alinan", ""),
                       "S2": S2, "O2": O2, "D2": D2, "rpn2": rpn2, "ap2": h["ap2"]})
        kp = []
        for r in db.query(ProsesKpSatiri).filter_by(adim_tanim_id=a.adim_tanim_id, urun_proses_id=u.id).order_by(ProsesKpSatiri.sira, ProsesKpSatiri.id).all():
            v = _j(r.valeo_json)
            kp.append({"surec_no": v.get("surec_no", ""), "surec_ad": v.get("surec_ad", ""), "ekipman": v.get("ekipman", ""), "no": v.get("no", str(r.sira)),
                       "urun": (v.get("urun", "") if v else (r.karakteristik or "")), "proses": v.get("proses", ""), "ozel": v.get("ozel", r.ozel or ""),
                       "sartname": v.get("sartname", ""), "teknik": v.get("teknik", r.olcum_araci or ""), "numune": v.get("numune", str(r.numune or "")), "siklik": v.get("siklik", r.siklik or ""),
                       "kontrol": v.get("kontrol", r.yontem or ""), "kayit": v.get("kayit", ""), "reaksiyon": v.get("reaksiyon", r.reaksiyon or ""), "sistem_satiri": not v})
        ops.append({"ad": katalog.get(a.adim_tanim_id, "?"), "vj": vj, "fmea": fm, "kp": kp, "sira": a.sira})
    return ops


def _baslik(u):
    v = _j(u.valeo_json)
    return v


def _urun_al(db, uid):
    u = db.get(UrunProsesi, uid)
    if not u:
        abort(404)
    return u


def _doc_no(s):
    m = re.search(r"T\d\s?P\s?[\d.]+\s?/\s?\d+", s or "")
    return m.group(0) if m else ""


# ─────────────────────────── HTML çıktılar ───────────────────────────
@app.route("/urun-proses/<int:uid>/cikti/akis")
@login_required
def valeo_akis(uid):
    db = get_db()
    try:
        u = _urun_al(db, uid)
        if request.args.get("excel") == "1":
            return _excel(db, u, "akis")
        v = _baslik(u)
        ops = [o for o in _veri(db, u) if o["vj"].get("etiket") or o["vj"].get("kutular")]
        flat = []
        for o in ops:
            kutular = o["vj"].get("kutular") or [{"metin": o["ad"], "sekil": 61}]
            for i, k in enumerate(kutular):
                flat.append({"etiket": (o["vj"].get("etiket") or o["ad"]) if i == 0 else "", "metin": k["metin"], "sekil": k.get("sekil"), "notlar": (o["vj"].get("notlar") or []) if i == 0 else [], "ilk": i == 0})
        log_action(db, "Çıktı", detay=f"Valeo Flow Chart: {u.urun_adi}")
        return render_template("valeo_akis.html", u=u, v=v, akis=v.get("akis", {}), satirlar=flat, doc_no=_doc_no(v.get("akis", {}).get("baslik", "")), bugun=date.today())
    finally:
        db.close()


@app.route("/urun-proses/<int:uid>/cikti/fmea")
@login_required
def valeo_fmea(uid):
    db = get_db()
    try:
        u = _urun_al(db, uid)
        if request.args.get("excel") == "1":
            return _excel(db, u, "fmea")
        v = _baslik(u)
        ops = [o for o in _veri(db, u) if o["fmea"]]
        log_action(db, "Çıktı", detay=f"Valeo PFMEA: {u.urun_adi}")
        return render_template("valeo_fmea.html", u=u, v=v, fm=v.get("fmea", {}), ops=ops, std=UP._std(u), standartlar=FC.STANDARTLAR, bugun=date.today())
    finally:
        db.close()


@app.route("/urun-proses/<int:uid>/cikti/kontrol-plani")
@login_required
def valeo_kp(uid):
    db = get_db()
    try:
        u = _urun_al(db, uid)
        if request.args.get("excel") == "1":
            return _excel(db, u, "kp")
        v = _baslik(u)
        ops = [o for o in _veri(db, u) if o["kp"]]
        log_action(db, "Çıktı", detay=f"Valeo Control Plan: {u.urun_adi}")
        return render_template("valeo_kp.html", u=u, v=v, cp=v.get("cp", {}), ops=ops, bugun=date.today())
    finally:
        db.close()


# ─────────────────────────── Excel ───────────────────────────
def _excel(db, u, tur):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    v = _baslik(u)
    ops = _veri(db, u)
    wb = Workbook(); wb.remove(wb.active)
    ince = Side(style="thin", color="555555"); kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    gri = PatternFill("solid", fgColor="D9E2F3"); bold = Font(bold=True)
    sar = Alignment(wrap_text=True, vertical="top")

    def antet(ws, baslik, kod, ncol):
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 42
        if os.path.isfile(LOGO):
            try:
                im = XImage(LOGO); im.width, im.height = 118, 40; ws.add_image(im, "A1")
            except Exception:  # noqa: BLE001
                pass
        ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=max(ncol - 2, 3))
        c = ws.cell(1, 2, baslik); c.font = Font(bold=True, size=14); c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=1, start_column=max(ncol - 1, 4), end_row=1, end_column=ncol)
        c = ws.cell(1, max(ncol - 1, 4), kod); c.alignment = Alignment(wrap_text=True, vertical="center"); c.font = Font(size=9)
        for col in range(1, ncol + 1):
            ws.cell(1, col).border = kenar

    if tur == "akis":
        ws = wb.create_sheet("Flow Chart")
        a = v.get("akis", {})
        antet(ws, a.get("baslik") or f"PRODUCTION FLOW CHART — {u.urun_adi}", f"Date: {a.get('tarih', '')}\nRev: {a.get('rev', '')}", 4)
        ws.append([]); ws.append(["Process", "Flow step", "Notes"])
        for c in ws[3]:
            c.font = bold; c.fill = gri; c.border = kenar
        for o in ops:
            if not (o["vj"].get("etiket") or o["vj"].get("kutular")):
                continue
            ws.append([o["vj"].get("etiket") or o["ad"], "\n↓\n".join(k["metin"] for k in o["vj"].get("kutular", [])), "\n".join(o["vj"].get("notlar", []))])
            for c in ws[ws.max_row]:
                c.alignment = sar; c.border = kenar
        for col, w in zip("ABCD", (38, 56, 38, 4)):
            ws.column_dimensions[col].width = w
        ad = "flow_chart"
    elif tur == "fmea":
        for o in [x for x in ops if x["fmea"]]:
            fb = o["vj"].get("fmea_baslik", {})
            ws = wb.create_sheet(re.sub(r"[\[\]\\/*?:]", "", o["ad"])[:28])
            antet(ws, "PROCESS FMEA", f"FMEA No: {fb.get('no', v.get('fmea', {}).get('no', ''))}\nRev. date: {fb.get('tarih_rev', '')}", 17)
            ws.append([f"Process: {fb.get('surec', o['ad'])}", "", "", f"Item: {fb.get('urun', u.urun_adi)}", "", "", f"Prepared by: {fb.get('hazirlayan', '')}", "", "", f"FMEA Date (Org.): {fb.get('tarih_org', '')}"])
            ws.append([f"Core Team: {fb.get('ekip', '')}"])
            ws.append(["Process Function", "Requirement", "Potential Failure Mode", "Potential Effects of Failure", "Severity", "Class", "Potential Causes of Failure", "Prevention", "Occur.", "Detection",
                       "Detect.", "RPN", "AP", "Recommended Actions", "Responsib. / Target Completion Date", "Actions Taken", "Sev. / Occ. / Det. / RPN"])
            for c in ws[ws.max_row]:
                c.font = bold; c.fill = gri; c.border = kenar; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            for r in o["fmea"]:
                ws.append([r["fn"], r["req"], r["mode"], r["eff"], r["S"], r["sinif"], r["cause"], r["prev"], r["O"], r["det"], r["D"], r["rpn"], FC.AP_AD.get(r["ap"], ""), r["act"], r["sorumlu"], r["alinan"],
                           " / ".join(str(x) for x in (r["S2"], r["O2"], r["D2"], r["rpn2"]) if x is not None)])
                for c in ws[ws.max_row]:
                    c.alignment = sar; c.border = kenar
            for col, w in zip("ABCDEFGHIJKLMNOPQ", (28, 28, 26, 30, 8, 7, 30, 30, 8, 30, 8, 8, 9, 24, 20, 22, 14)):
                ws.column_dimensions[col].width = w
            ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = 9; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        ad = "pfmea"
    else:
        cp = v.get("cp", {})
        ws = wb.create_sheet("Control Plan")
        antet(ws, "CONTROL PLAN", f"Control Plan No: {cp.get('no', '')}\nItem Ref. No: {cp.get('ref', u.urun_kodu or '')}", 14)
        ws.append([f"Item Name: {cp.get('ad', u.urun_adi)}", "", "", f"Key Contact: {cp.get('iletisim', '')}", "", "", "", f"Core Team: {cp.get('ekip', '')}"])
        ws.append([f"Supplier: DEBAK DENİZLİ BAGALİT KALIP SAN. & TİC. A.Ş.   Supplier Code: {cp.get('kod', '')}", "", "", f"Date (Orj.): {cp.get('tarih_org', '')}", "", "", "", "Date (Rev.): " + " · ".join(cp.get("revler", []))])
        ws.append(["Process No", "Process Description", "Machine / Device / Production Equipment, Material", "Char. No", "Product", "Process", "Special Char. Class", "Product / Process Specification / Tolerance",
                   "Evaluation / Measurement Technique", "Sample Size", "Frequency", "Control Method", "Records", "Reaction Plan"])
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = gri; c.border = kenar; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        son = None
        for o in [x for x in ops if x["kp"]]:
            for r in o["kp"]:
                anahtar = (r["surec_no"], r["surec_ad"])
                ws.append([r["surec_no"] if anahtar != son else "", r["surec_ad"] if anahtar != son else "", r["ekipman"], r["no"], r["urun"], r["proses"], r["ozel"], r["sartname"], r["teknik"], r["numune"], r["siklik"],
                           r["kontrol"], r["kayit"], r["reaksiyon"]])
                son = anahtar
                for c in ws[ws.max_row]:
                    c.alignment = sar; c.border = kenar
        ws.append([]); ws.append(["Reaction Plan Definitions"]); ws.cell(ws.max_row, 1).font = bold
        for t in cp.get("reaksiyon_tablo", []):
            ws.append([t["no"], t["metin"]])
            ws.cell(ws.max_row, 2).alignment = sar
        for col, w in zip("ABCDEFGHIJKLMN", (9, 26, 28, 8, 28, 28, 8, 28, 28, 12, 16, 30, 30, 10)):
            ws.column_dimensions[col].width = w
        ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = 9; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ad = "control_plan"
    if not wb.worksheets:
        wb.create_sheet("Boş").append(["Bu ürün için veri yok."])
    b = io.BytesIO(); wb.save(b); b.seek(0)
    log_action(db, "Çıktı", detay=f"Valeo Excel ({ad}): {u.urun_adi}")
    return send_file(b, as_attachment=True, download_name=f"{ad}_{u.urun_kodu or u.id}_{date.today():%Y%m%d}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
