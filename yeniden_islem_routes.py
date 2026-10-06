"""
DYS — Yeniden İşlem Takip (2026): uygunsuz / karantinaya alınan ürünün yeniden işlem (ayıklama, çapak alma, %100 kontrol …) kaydı, sonucu ve maliyeti.
Eski D06 F06 'Uygun Olmayan Ürün & Yeniden İşlem Takip Listesi' Excel'inin yerini alır (Kalite menüsü).
  /kalite/yeniden-islem            liste + filtre (yıl / statü / durum / ara) + özet kartları + Excel
  /kalite/yeniden-islem/<id>       kayıt / düzenleme (sonuç ve maliyet otomatik)
  /kalite/yeniden-islem/rapor      yıl ve ürün bazında özet (miktar, hatalı, maliyet)
  /kalite/yeniden-islem/df         düzeltici faaliyet notları
Maliyet: imha = hatalı miktar × parça birim maliyeti; işçilik = süre (saat) × operatör sayısı × saatlik ücret (Sistem ayarı 'yeniden_islem_saat_ucreti', öntanımlı 5);
toplam = imha + işçilik. Elle girilen maliyet alanları öncelikli. Yetki: kayıt Sadece Görüntüleme hariç; silme Admin / Doküman Kontrol.
"""
import io
from datetime import date, datetime

from flask import render_template, request, redirect, url_for, flash, abort, session, send_file

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import YenidenIslemKaydi, YenidenIslemDF, SistemAyar  # noqa: E402

DURUMLAR = ("Açık", "Devam ediyor", "Tamamlandı")
SAYFA = 50


def _yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _yazabilir():
    return session.get("rol") != "Sadece Görüntüleme"


def _k(s):
    return " ".join((s or "").replace("İ", "i").replace("I", "ı").lower().split())


def _sayi(v):
    v = (str(v) if v is not None else "").strip().replace(",", ".")
    if not v:
        return None
    try:
        return float(v)
    except ValueError:
        raise ValueError(v)


def _tarih(v):
    try:
        return datetime.strptime(v, "%Y-%m-%d").date() if v else None
    except ValueError:
        return None


def _saat_ucreti(db):
    a = db.query(SistemAyar).filter_by(anahtar="yeniden_islem_saat_ucreti").first()
    try:
        return float(a.deger) if a and a.deger else 5.0
    except ValueError:
        return 5.0


def fmt(v):
    """Binlik ayraçlı tam sayı gösterimi (1.234)."""
    try:
        return "{:,.0f}".format(v or 0).replace(",", ".")
    except (TypeError, ValueError):
        return "—"


def durum_goster(r):
    if r.durum:
        return r.durum
    return "Tamamlandı" if (r.islem_tarihi or r.saglam_miktar is not None) else "Açık"


def _maliyet(db, r):
    """Boş maliyet alanlarını hesaplar (elle girilenlere dokunmaz)."""
    if r.imha_maliyeti is None and r.hatali_miktar is not None and r.parca_birim_maliyet is not None:
        r.imha_maliyeti = round(r.hatali_miktar * r.parca_birim_maliyet, 4)
    if r.iscilik_maliyeti is None and r.sure is not None:
        saat = r.sure / 60 if (r.sure_birim or "dk").lower().startswith("d") else r.sure
        r.iscilik_maliyeti = round(saat * (r.operator_sayisi or 1) * _saat_ucreti(db), 4)
    if r.toplam_maliyet is None and (r.imha_maliyeti is not None or r.iscilik_maliyeti is not None):
        r.toplam_maliyet = round((r.imha_maliyeti or 0) + (r.iscilik_maliyeti or 0), 4)


def _filtre(db, args):
    q = db.query(YenidenIslemKaydi)
    yil = args.get("yil", type=int)
    if yil:
        q = q.filter(YenidenIslemKaydi.giris_tarihi >= date(yil, 1, 1), YenidenIslemKaydi.giris_tarihi <= date(yil, 12, 31))
    if args.get("statu"):
        q = q.filter(YenidenIslemKaydi.statu == args.get("statu"))
    L = q.order_by(YenidenIslemKaydi.giris_tarihi.desc(), YenidenIslemKaydi.id.desc()).all()
    d = args.get("durum", "")
    if d:
        L = [r for r in L if durum_goster(r) == d]
    ara = _k(args.get("q"))
    if ara:
        L = [r for r in L if ara in _k(f"{r.urun} {r.stok_adi} {r.stok_kodu} {r.is_emri} {r.lot_no} {r.uygunsuzluk} {r.yapilan_islem} {r.sorumlu}")]
    return L


def _ozet(L):
    return {"adet": len(L), "miktar": sum(r.miktar or 0 for r in L), "hatali": sum(r.hatali_miktar or 0 for r in L), "saglam": sum(r.saglam_miktar or 0 for r in L),
            "maliyet": sum(r.toplam_maliyet or 0 for r in L), "acik": sum(1 for r in L if durum_goster(r) != "Tamamlandı")}


@app.route("/kalite/yeniden-islem", methods=["GET", "POST"])
@login_required
def yeniden_islem_liste():
    db = get_db()
    try:
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            r = YenidenIslemKaydi(kaynak="elle", olusturan_id=session.get("user_id"), giris_tarihi=date.today(), durum="Açık")
            err = _doldur(db, r, request.form)
            if err:
                flash(err, "error")
                return redirect(url_for("yeniden_islem_liste"))
            db.add(r); db.commit()
            log_action(db, "Oluşturma", detay=f"Yeniden işlem kaydı: {r.urun} — {r.miktar}")
            return redirect(url_for("yeniden_islem_kayit", rid=r.id))
        args = request.args
        L = _filtre(db, args)
        if args.get("excel") == "1":
            return _excel(L)
        sayfa = max(args.get("s", 1, type=int), 1)
        toplam = len(L)
        yillar = sorted({d.year for (d,) in db.query(YenidenIslemKaydi.giris_tarihi).all() if d}, reverse=True)
        statuler = sorted({s for (s,) in db.query(YenidenIslemKaydi.statu).distinct().all() if s}, key=_k)
        return render_template("yeniden_islem_liste.html", liste=L[(sayfa - 1) * SAYFA: sayfa * SAYFA], ozet=_ozet(L), toplam=toplam, sayfa=sayfa, sayfa_sayisi=max((toplam + SAYFA - 1) // SAYFA, 1),
                               yillar=yillar, statuler=statuler, args=args, durumlar=DURUMLAR, durum_goster=durum_goster, fmt=fmt, yazabilir=_yazabilir())
    finally:
        db.close()


def _doldur(db, r, f):
    try:
        for alan in ("miktar", "saglam_miktar", "hatali_miktar", "sure", "operator_sayisi", "birim_fiyat", "parca_birim_maliyet", "imha_maliyeti", "geri_kazanim", "iscilik_maliyeti", "toplam_maliyet"):
            if alan in f:
                setattr(r, alan, _sayi(f.get(alan)))
    except ValueError as e:
        return f"Sayı olmayan değer: {e}"
    for alan, n in (("saha", 80), ("statu", 60), ("urun", 200), ("stok_adi", 250), ("stok_kodu", 60), ("is_emri", 60), ("lot_no", 100), ("uygunsuzluk", 1000), ("yapilan_islem", 1000),
                    ("aksiyon", 1000), ("sorumlu", 100), ("sure_birim", 10), ("doviz", 10)):
        if alan in f:
            setattr(r, alan, (f.get(alan) or "").strip()[:n] or None)
    for alan in ("giris_tarihi", "islem_tarihi", "termin"):
        if alan in f:
            setattr(r, alan, _tarih(f.get(alan)))
    if f.get("durum") in DURUMLAR:
        r.durum = f.get("durum")
    if not (r.urun or r.stok_adi):
        return "Ürün veya stok adı zorunludur."
    if r.miktar is None:
        return "Miktar zorunludur."
    _maliyet(db, r)
    return None


@app.route("/kalite/yeniden-islem/<int:rid>", methods=["GET", "POST"])
@login_required
def yeniden_islem_kayit(rid):
    db = get_db()
    try:
        r = db.get(YenidenIslemKaydi, rid)
        if not r:
            abort(404)
        if request.method == "POST":
            if request.form.get("action") == "sil":
                if not _yonetici():
                    abort(403)
                db.delete(r); db.commit(); flash("Kayıt silindi.", "success")
                return redirect(url_for("yeniden_islem_liste"))
            if not _yazabilir():
                abort(403)
            if request.form.get("maliyet_sifirla") == "1":
                r.imha_maliyeti = r.iscilik_maliyeti = r.toplam_maliyet = None
            err = _doldur(db, r, request.form)
            if err:
                flash(err, "error")
            else:
                if request.form.get("tamamla") == "1":
                    r.durum = "Tamamlandı"
                    r.islem_tarihi = r.islem_tarihi or date.today()
                db.commit(); flash("Kaydedildi.", "success")
                log_action(db, "Güncelleme", detay=f"Yeniden işlem kaydı güncellendi: #{r.id} {r.urun}")
            return redirect(url_for("yeniden_islem_kayit", rid=r.id))
        return render_template("yeniden_islem_kayit.html", r=r, durumlar=DURUMLAR, durum=durum_goster(r), yazabilir=_yazabilir(), yonetici=_yonetici(), saat_ucreti=_saat_ucreti(db))
    finally:
        db.close()


@app.route("/kalite/yeniden-islem/rapor")
@login_required
def yeniden_islem_rapor():
    db = get_db()
    try:
        L = db.query(YenidenIslemKaydi).all()
        yil, urun = {}, {}
        for r in L:
            y = r.giris_tarihi.year if r.giris_tarihi else 0
            for d, k in ((yil, y), (urun, (r.urun or r.stok_adi or "—").strip())):
                a = d.setdefault(k, {"adet": 0, "miktar": 0.0, "hatali": 0.0, "maliyet": 0.0})
                a["adet"] += 1; a["miktar"] += r.miktar or 0; a["hatali"] += r.hatali_miktar or 0; a["maliyet"] += r.toplam_maliyet or 0
        statu = {}
        for r in L:
            a = statu.setdefault((r.statu or "—").strip(), {"adet": 0, "miktar": 0.0, "maliyet": 0.0})
            a["adet"] += 1; a["miktar"] += r.miktar or 0; a["maliyet"] += r.toplam_maliyet or 0
        top_urun = sorted(urun.items(), key=lambda x: -x[1]["maliyet"])[:20]
        return render_template("yeniden_islem_rapor.html", yil=sorted(yil.items(), reverse=True), top_urun=top_urun, statu=sorted(statu.items(), key=lambda x: -x[1]["adet"]), ozet=_ozet(L), fmt=fmt)
    finally:
        db.close()


@app.route("/kalite/yeniden-islem/df", methods=["GET", "POST"])
@login_required
def yeniden_islem_df():
    db = get_db()
    try:
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            f = request.form
            if f.get("action") == "sil":
                if not _yonetici():
                    abort(403)
                x = db.get(YenidenIslemDF, f.get("id", type=int) or 0)
                if x:
                    db.delete(x); db.commit()
                return redirect(url_for("yeniden_islem_df"))
            t = _tarih(f.get("tarih_d"))
            x = YenidenIslemDF(tarih=(t.strftime("%d.%m.%Y") if t else None), tarih_d=t, urun=(f.get("urun") or "").strip()[:200] or None, uygunsuzluk=(f.get("uygunsuzluk") or "").strip()[:1000] or None,
                               acil_onlem=(f.get("acil_onlem") or "").strip()[:1000] or None, kok_neden=(f.get("kok_neden") or "").strip()[:1000] or None, planlanan=(f.get("planlanan") or "").strip()[:1000] or None,
                               sorumlu=(f.get("sorumlu") or "").strip()[:150] or None, hedef_tarih=_tarih(f.get("hedef_tarih")), gerceklesen=(f.get("gerceklesen") or "").strip()[:1000] or None,
                               gerceklesen_tarih=_tarih(f.get("gerceklesen_tarih")), kaynak="elle")
            if not (x.urun or x.uygunsuzluk):
                flash("Ürün veya uygunsuzluk yazın.", "error")
            else:
                db.add(x); db.commit(); flash("Eklendi.", "success")
            return redirect(url_for("yeniden_islem_df"))
        L = db.query(YenidenIslemDF).order_by(YenidenIslemDF.tarih_d.desc(), YenidenIslemDF.id.desc()).all()
        return render_template("yeniden_islem_df.html", liste=L, yazabilir=_yazabilir(), yonetici=_yonetici())
    finally:
        db.close()


def _excel(L):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Yeniden İşlem"
    ws.append(["Giriş tarihi", "Saha", "Statü", "Ürün", "Stok adı", "Stok kodu", "İş emri", "Lot", "Miktar", "Uygunsuzluk", "İşlem tarihi", "Yapılan işlem", "Aksiyon", "Sorumlu", "Termin", "Durum",
               "Sağlam", "Hatalı", "Süre", "Birim", "Operatör", "Birim fiyat", "Döviz", "Parça maliyeti", "İmha maliyeti", "Geri kazanım", "İşçilik", "Toplam maliyet"])
    for r in L:
        ws.append([r.giris_tarihi, r.saha, r.statu, r.urun, r.stok_adi, r.stok_kodu, r.is_emri, r.lot_no, r.miktar, r.uygunsuzluk, r.islem_tarihi, r.yapilan_islem, r.aksiyon, r.sorumlu, r.termin,
                   durum_goster(r), r.saglam_miktar, r.hatali_miktar, r.sure, r.sure_birim, r.operator_sayisi, r.birim_fiyat, r.doviz, r.parca_birim_maliyet, r.imha_maliyeti, r.geri_kazanim,
                   r.iscilik_maliyeti, r.toplam_maliyet])
    for c in ws["A"][1:] + ws["K"][1:] + ws["O"][1:]:
        c.number_format = "DD.MM.YYYY"
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return send_file(b, as_attachment=True, download_name="yeniden_islem_takip.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
