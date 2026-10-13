"""
DYS — Prosedür DOCX şablonu (Debak antet / alt bilgi)
=====================================================
Referans prosedür düzenine uygun Word belgesi üretir:
- Antet: logo + süreç adı + doküman başlığı + kod / rev / tarih / sayfa
- Gövde: 1.0–7.0 bölümleri (DOKUMAN_SABLONLARI Prosedür)
- Alt bilgi: HAZIRLAYAN / KONTROL EDEN / ONAYLAYAN

Not (12.10.2026): Kaynak kod kayıptı; yalnız Python 3.10 ile derlenmiş procedure_docx.pyc vardı
ve Python 3.13 bunu yükleyemediği için tüm prosedür sayfaları 500 veriyordu. Bu dosya o derlemeden
birebir yeniden kuruldu (aynı fonksiyonlar, sabitler ve davranış).
"""

from __future__ import annotations

import json
import os
from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: F401
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, Twips  # noqa: F401

_BRANDING_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static", "branding")
LOGO_PATH = os.path.join(_BRANDING_DIR, "debak-logo.png")

# (icerik anahtarı, başlık, yer tutucu)
PROSEDUR_BOLUMLERI = (
    ("amac", "1.0\tAMAÇ", "Prosedürün amacını yazın (hangi ihtiyacı karşıladığı)."),
    ("uygulama_alani", "2.0\tUYGULAMA ALANI", "Kapsanan süreç, faaliyet, ürün ve sınırları yazın."),
    ("tanimlar", "3.0\tTANIMLAR", "Terim ve kısaltmaları numaralandırarak tanımlayın (ör. 3.1 PPAP: …)."),
    ("sorumluluklar", "4.0\tSORUMLULUKLAR", "Rolleri ve sorumluluklarını madde madde yazın."),
    ("uygulama_sekli", "5.0\tUYGULAMA ŞEKLİ",
     "Uygulama adımlarını alt başlıklar (A, B, C…) ve numaralı maddelerle yazın."),
    ("kayitlar", "6.0\tKAYITLARIN MUHAFAZASI", "Doküman listesinden arayıp seçin (kayıt formları vb.)."),
    ("ilgili_dokumanlar", "7.0\tİLGİLİ DOKÜMANLAR", "Doküman listesinden arayıp seçin."),
)


def _set_run_font(run, *, name="Cambria Math", size_pt=9, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def _set_cell_border(cell, **kwargs):
    """kwargs: top/left/bottom/right = 'sz', 'val', 'color'."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)
        if edge_data is None:
            continue
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), edge_data.get("val", "single"))
        element.set(qn("w:sz"), str(edge_data.get("sz", 4)))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), edge_data.get("color", "000000"))
        tcBorders.append(element)
    tcPr.append(tcBorders)


def _shade_cell(cell, fill="FFFFFF"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def _set_cell_vertical_align(cell, align="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    v_align = tcPr.find(qn("w:vAlign"))
    if v_align is None:
        v_align = OxmlElement("w:vAlign")
        tcPr.append(v_align)
    v_align.set(qn("w:val"), align)


def _merge_cells_horizontal(row, start, end):
    cell = row.cells[start]
    for i in range(start + 1, end + 1):
        cell.merge(row.cells[i])
    return cell


ANTET_SHADE = "D9E2F3"


def _clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            p.remove(child)


def _add_page_number_field(paragraph):
    """Sayfa X / Y alanları (Word PAGE / NUMPAGES)."""
    run = paragraph.add_run()
    _set_run_font(run, size_pt=9)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)

    run2 = paragraph.add_run(" / ")
    _set_run_font(run2, size_pt=9)

    run3 = paragraph.add_run()
    _set_run_font(run3, size_pt=9)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " NUMPAGES "
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run3._r.append(fld_char_begin)
    run3._r.append(instr)
    run3._r.append(fld_char_sep)
    run3._r.append(fld_char_end)


def _write_cell(cell, text, *, size_pt=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, clear=True):
    if clear:
        cell.text = ""
        p = cell.paragraphs[0]
    else:
        p = cell.add_paragraph()
    p.alignment = align
    run = p.add_run(text or "")
    _set_run_font(run, size_pt=size_pt, bold=bold)
    return p


def _build_header(section, meta):
    header = section.header
    header.is_linked_to_previous = False

    for p in header.paragraphs:
        _clear_paragraph(p)

    table = header.add_table(rows=4, cols=4, width=Cm(18.2))
    table.autofit = False
    widths = (Cm(3.2), Cm(8.5), Cm(3.0), Cm(3.5))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]

    border = {"sz": 4, "val": "single", "color": "000000"}
    for row in table.rows:
        for cell in row.cells:
            _set_cell_border(cell, top=border, left=border, bottom=border, right=border)

    # Logo (sol, iki satır)
    logo_cell = table.cell(0, 0)
    logo_cell.merge(table.cell(1, 0))
    logo_cell.text = ""
    _set_cell_vertical_align(logo_cell, "center")
    p = logo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.isfile(LOGO_PATH):
        run = p.add_run()
        run.add_picture(LOGO_PATH, width=Cm(2.6))

    # Süreç adı (orta, iki satır)
    process_cell = table.cell(0, 1)
    process_cell.merge(table.cell(1, 1))
    _write_cell(
        process_cell,
        meta.get("surec_ust") or "",
        size_pt=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _shade_cell(process_cell, ANTET_SHADE)
    _set_cell_vertical_align(process_cell, "center")

    meta_specs = (
        (0, 2, "Doküman Kodu", None),
        (0, 3, meta.get("dokuman_no") or "", WD_ALIGN_PARAGRAPH.CENTER),
        (1, 2, "Rev. No", None),
        (1, 3, str(meta.get("revizyon_no", 0)), WD_ALIGN_PARAGRAPH.CENTER),
        (2, 2, "Rev. Tarihi", None),
        (2, 3, meta.get("rev_tarihi") or "", WD_ALIGN_PARAGRAPH.CENTER),
        (3, 2, "Sayfa", None),
    )
    for row_i, col_i, text, align in meta_specs:
        cell = table.cell(row_i, col_i)
        _write_cell(
            cell,
            text,
            size_pt=9,
            bold=True,
            align=align or WD_ALIGN_PARAGRAPH.LEFT,
        )
        _shade_cell(cell, ANTET_SHADE)
        _set_cell_vertical_align(cell, "center")

    # Doküman başlığı (alt iki satır, sol iki sütun)
    title_top = _merge_cells_horizontal(table.rows[2], 0, 1)
    title_bot = _merge_cells_horizontal(table.rows[3], 0, 1)
    title_cell = title_top.merge(title_bot)
    _write_cell(
        title_cell,
        meta.get("baslik") or "",
        size_pt=12,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _set_cell_vertical_align(title_cell, "center")

    page_cell = table.cell(3, 3)
    page_cell.text = ""
    _shade_cell(page_cell, ANTET_SHADE)
    _set_cell_vertical_align(page_cell, "center")
    pp = page_cell.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number_field(pp)


def _build_footer(section, meta):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        _clear_paragraph(p)

    table = footer.add_table(rows=2, cols=3, width=Cm(18.2))
    border = {"sz": 4, "val": "single", "color": "000000"}
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(6.0)
            _set_cell_border(cell, top=border, left=border, bottom=border, right=border)

    labels = ("HAZIRLAYAN", "KONTROL EDEN", "ONAYLAYAN")
    values = (
        meta.get("hazirlayan") or "",
        meta.get("kontrol_eden") or "",
        meta.get("onaylayan") or "",
    )
    for i, label in enumerate(labels):
        _write_cell(table.cell(0, i), label, size_pt=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade_cell(table.cell(0, i), "F3F4F6")
        _write_cell(table.cell(1, i), values[i], size_pt=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def format_procedure_refs(value):
    """icerik kayitlar/ilgili_dokumanlar değerini metin satırlarına çevirir.

    Desteklenen biçimler:
    - JSON liste: [{"dokuman_no": "...", "baslik": "..."}, ...]
    - Düz metin (eski kayıtlar)
    """
    if value is None:
        return ""
    if isinstance(value, list):
        items = value
    else:
        text = str(value).strip()
        if not text:
            return ""
        try:
            parsed = json.loads(text)
            items = parsed if isinstance(parsed, list) else None
        except (TypeError, ValueError, json.JSONDecodeError):
            items = None
        if items is None:
            return text

    lines = []
    for item in items:
        if isinstance(item, dict):
            no = (item.get("dokuman_no") or "").strip()
            baslik = (item.get("baslik") or "").strip()
            if no and baslik:
                lines.append(f"{no} — {baslik}")
            elif no or baslik:
                lines.append(no or baslik)
        else:
            s = str(item).strip()
            if s:
                lines.append(s)
    return "\n".join(lines)


def parse_procedure_refs(value):
    """Görüntüleme için yapılandırılmış liste döndürür."""
    if not value:
        return []
    if isinstance(value, list):
        return [i for i in value if isinstance(i, dict)]
    text = str(value).strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return [i for i in parsed if isinstance(i, dict)]
    except (TypeError, ValueError, json.JSONDecodeError):
        pass
    return [
        {"dokuman_no": "", "baslik": line.strip(), "id": None}
        for line in text.splitlines() if line.strip()
    ]


def _parse_icerik(icerik):
    if not icerik:
        return {}
    if isinstance(icerik, dict):
        return icerik
    try:
        return json.loads(icerik)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}


def build_procedure_docx(meta, icerik=None) -> BytesIO:
    """Referans düzeninde prosedür DOCX üretir; BytesIO döner."""
    icerik = _parse_icerik(icerik)
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.25)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(1.5)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)

    _build_header(section, meta)
    _build_footer(section, meta)

    ref_keys = {"kayitlar", "ilgili_dokumanlar"}
    for key, baslik, _ph in PROSEDUR_BOLUMLERI:
        heading = doc.add_paragraph()
        heading.style = doc.styles["Heading 1"]
        run = heading.add_run(baslik.replace("\t", "  "))
        _set_run_font(run, name="Cambria Math", size_pt=12, bold=True)

        raw = icerik.get(key)
        if key in ref_keys:
            body_text = format_procedure_refs(raw).strip()
        else:
            body_text = (raw or "").strip() if isinstance(raw, str) else format_procedure_refs(raw).strip()
        if not body_text:
            body_text = "…"
        for line in body_text.splitlines() or ["…"]:
            p = doc.add_paragraph()
            prefix = "• " if key in ref_keys and line != "…" else ""
            run = p.add_run(prefix + line)
            _set_run_font(run, name="Cambria Math", size_pt=11, bold=False)
            p.paragraph_format.space_after = Pt(6)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def meta_from_document(doc, db=None):
    """SQLAlchemy Document nesnesinden antet/alt bilgi meta sözlüğü."""
    surec = getattr(doc, "surec", None)
    surec_ust = ""
    if surec is not None:
        surec_ust = f"{surec.kod or ''} {(surec.ad or '').upper()}".strip()

    rev_tarihi = ""
    if getattr(doc, "yururluk_tarihi", None):
        rev_tarihi = doc.yururluk_tarihi.strftime("%d.%m.%Y")
    elif getattr(doc, "guncelleme_tarihi", None):
        dt = doc.guncelleme_tarihi
        rev_tarihi = dt.strftime("%d.%m.%Y") if isinstance(dt, datetime) else str(dt)

    def _role_name(user_rel, fallback=""):
        if user_rel is not None:
            return (user_rel.unvan or user_rel.ad_soyad or "").upper()
        return (fallback or "").upper()

    return {
        "surec_ust": surec_ust,
        "baslik": doc.baslik or "",
        "dokuman_no": doc.dokuman_no or "",
        "revizyon_no": doc.revizyon_no if doc.revizyon_no is not None else 0,
        "rev_tarihi": rev_tarihi or date.today().strftime("%d.%m.%Y"),
        "hazirlayan": _role_name(getattr(doc, "hazirlayan", None)),
        "kontrol_eden": _role_name(getattr(doc, "kontrol_eden", None)),
        "onaylayan": _role_name(getattr(doc, "onaylayan", None)),
    }


def save_procedure_docx(doc, target_path):
    """Prosedür DOCX dosyasını diske yazar."""
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    buf = build_procedure_docx(meta_from_document(doc), doc.icerik_json)
    with open(target_path, "wb") as f:
        f.write(buf.getvalue())
    return target_path
