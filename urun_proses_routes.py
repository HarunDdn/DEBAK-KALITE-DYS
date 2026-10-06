"""
DYS — Ürün Prosesi (2026): müşteri → ürün → proses akışı (adımlar) → FMEA + kontrol planı otomatik → yalnız ÖLÇÜ girişi.
  /urun-proses                      ürünler (müşteri bazlı), yeni ürün
  /urun-proses/kutuphane            proses adımı kataloğu + adım başına FMEA şablonu (P-FMEA, GSI-RD-370) ve kontrol planı şablonu
  /urun-proses/<id>                 ürünün proses akışı: adım ekle / sırala; her adımda otomatik gelen FMEA (genel + ürüne özgü, AP otomatik) ve
                                    kontrol planı (şablondan kopyalanır; ürüne özgü nominal / tolerans / ölçüm aracı girilir)
  /urun-proses/<id>/olcum           ölçü girişi: yalnız değer girilir, OK / NOK limitlerden otomatik; numune sayısı kontrol planından
  /urun-proses/<id>/olcumler        ölçü geçmişi, istatistik (ort., min, maks, NOK %, Cp / Cpk) ve Excel
Akış: adım eklenince → o adımın FMEA şablon satırları ürüne otomatik görünür (genel satırlar salt okunur, ürüne özgü satır eklenebilir), kontrol planı şablon
satırları ürüne kopyalanır; tolerans girilince ölçü girişinde görünür. FMEA şablonunda sonradan eklenen satır tüm ürünlerde görünür; kontrol planı için
'Şablondan eksikleri getir'. Yetki: tanım (katalog, şablon, ürün prosesi, FMEA / kontrol planı satırları) Admin / Doküman Kontrol; ölçü girişi Sadece Görüntüleme hariç.
AP yalnız matristen (fmea_core); RPN yalnız bilgi.
"""
import io
import statistics
from datetime import date, datetime

from flask import render_template, request, redirect, url_for, flash, abort, session, jsonify, send_file

from app_runtime import host as _host
_h = _host()
app = _h.app
login_required = _h.login_required
get_db = _h.get_db
log_action = _h.log_action

from models import (ProsesAdimTanim, UrunProsesi, UrunProsesAdimi, ProsesFmeaSatiri, ProsesKpSatiri, OlcumKaydi, User)  # noqa: E402
import fmea_core as FC  # noqa: E402

VARSAYILAN_ADIMLAR = ("Enjeksiyon Hazırlama", "Enjeksiyon", "Çapak Alma", "Fırın", "Kalite Kontrol")
FMEA_METIN = ("fn", "eff", "mode", "cause", "prev", "det", "act")
FMEA_PUAN = ("S", "O", "D", "O2", "D2")
KP_METIN = {"karakteristik": 200, "tur": 12, "birim": 20, "ozel": 4, "olcum_araci": 150, "siklik": 100, "yontem": 300, "reaksiyon": 400}
KP_SAYI = ("nominal", "alt_limit", "ust_limit")
TURLER = ("Ölçü", "Görsel")


def _yonetici():
    return session.get("rol") in ("Admin", "Doküman Kontrol")


def _yazabilir():
    return session.get("rol") != "Sadece Görüntüleme"


def _k(s):
    return " ".join((s or "").replace("İ", "i").replace("I", "ı").lower().split())


def _katalog(db):
    if not db.query(ProsesAdimTanim).count():
        for i, ad in enumerate(VARSAYILAN_ADIMLAR, 1):
            db.add(ProsesAdimTanim(ad=ad, sira=i, aktif=True))
        db.commit()
    return db.query(ProsesAdimTanim).order_by(ProsesAdimTanim.sira, ProsesAdimTanim.id).all()


def _sayi(v):
    v = (str(v) if v is not None else "").strip().replace(",", ".")
    if not v:
        return None
    try:
        return float(v)
    except ValueError:
        raise ValueError("Sayı girin")


def _std(x):
    """Ürün / FMEA satırı için standart kodu: 'GSI' (Valeo GSI-RD-370, varsayılan) | 'VDA' (AIAG-VDA)."""
    v = x if isinstance(x, str) or x is None else (getattr(x, "fmea_standart", None) or getattr(x, "standart", None))
    return v if v in FC.STANDARTLAR else "GSI"


def _std_sorgu(q, std):
    """Standarda göre FMEA satırı filtresi (boş = GSI)."""
    if std == "VDA":
        return q.filter(ProsesFmeaSatiri.standart == "VDA")
    return q.filter((ProsesFmeaSatiri.standart.is_(None)) | (ProsesFmeaSatiri.standart == "GSI"))


def _std_varsayilan(musteri):
    return "GSI" if "valeo" in _k(musteri) else "VDA"


def _fmea_json(r):
    d = {k: getattr(r, k) for k in FMEA_METIN + FMEA_PUAN}
    h = FC.hesapla(d, "P", _std(r.standart))
    return {"id": r.id, **{k: ("" if v is None else v) for k, v in d.items()}, "ap": h["ap"], "ap_ad": FC.AP_AD[h["ap"]], "ap2": h["ap2"],
            "ap2_ad": FC.AP_AD[h["ap2"]], "rpn": h["rpn"] if h["rpn"] is not None else "", "uyari": h["uyari"]}


def _kp_json(r):
    return {"id": r.id, **{k: ("" if getattr(r, k) is None else getattr(r, k)) for k in list(KP_METIN) + list(KP_SAYI) + ["numune", "fmea_satir_id", "kaynak_id"]}}


# ─────────────────────────── ürünler ───────────────────────────
@app.route("/urun-proses", methods=["GET", "POST"])
@login_required
def urun_proses_liste():
    db = get_db()
    try:
        _katalog(db)
        if request.method == "POST":
            if not _yonetici():
                abort(403)
            f = request.form
            musteri, urun = " ".join((f.get("musteri") or "").split()), (f.get("urun_adi") or "").strip()
            if not (musteri and urun):
                flash("Müşteri ve ürün adı zorunludur.", "error")
                return redirect(url_for("urun_proses_liste"))
            for (m,) in db.query(UrunProsesi.musteri).distinct().all():       # aynı müşterinin farklı yazımı
                if _k(m) == _k(musteri):
                    musteri = m
                    break
            u = UrunProsesi(musteri=musteri[:150], urun_adi=urun[:200], urun_kodu=(f.get("urun_kodu") or "").strip()[:80] or None,
                            teknik_resim_id=f.get("teknik_resim_id", type=int),
                            fmea_standart=f.get("fmea_standart") if f.get("fmea_standart") in FC.STANDARTLAR else _std_varsayilan(musteri), aciklama=(f.get("aciklama") or "").strip()[:500] or None,
                            aktif=True, olusturan_id=session.get("user_id"))
            db.add(u); db.commit()
            log_action(db, "Oluşturma", detay=f"Ürün prosesi: {u.musteri} / {u.urun_adi}")
            return redirect(url_for("urun_proses_agac", uid=u.id))
        musteri = (request.args.get("musteri") or "").strip()
        q = (request.args.get("q") or "").strip().lower()
        L = [u for u in db.query(UrunProsesi).order_by(UrunProsesi.musteri, UrunProsesi.urun_adi).all()
             if u.aktif is not False and (not musteri or _k(u.musteri) == _k(musteri)) and (not q or q in _k(f"{u.urun_adi} {u.urun_kodu} {u.musteri}"))]
        adim_sayi = {}
        for a in db.query(UrunProsesAdimi).all():
            adim_sayi[a.urun_proses_id] = adim_sayi.get(a.urun_proses_id, 0) + 1
        olc = {}
        for (up, n) in db.query(OlcumKaydi.urun_proses_id, OlcumKaydi.id).all():
            olc[up] = olc.get(up, 0) + 1
        musteriler = sorted({u.musteri for u in db.query(UrunProsesi).all()}, key=_k)
        teknik = []
        try:
            from models import TeknikResim
            teknik = [(t.id, f"{t.musteri} — {t.urun_adi} ({t.resim_no})") for t in db.query(TeknikResim).filter(TeknikResim.aktif.isnot(False)).order_by(TeknikResim.musteri, TeknikResim.urun_adi).all()]
        except Exception:  # noqa: BLE001
            pass
        return render_template("urun_proses_liste.html", liste=L, adim_sayi=adim_sayi, olc=olc, musteriler=musteriler, musteri=musteri, q=request.args.get("q", ""),
                               yonetici=_yonetici(), teknik=teknik)
    finally:
        db.close()


# ─────────────────────────── kütüphane (şablonlar) ───────────────────────────
@app.route("/urun-proses/kutuphane", methods=["GET", "POST"])
@login_required
def urun_proses_kutuphane():
    db = get_db()
    try:
        katalog = _katalog(db)
        if request.method == "POST":
            if not _yonetici():
                abort(403)
            f = request.form
            a = f.get("action")
            if a == "adim_ekle":
                ad = " ".join((f.get("ad") or "").split())
                if not ad or any(_k(x.ad) == _k(ad) for x in katalog):
                    flash("Adım adı boş ya da zaten var.", "error")
                else:
                    db.add(ProsesAdimTanim(ad=ad[:120], sira=max([x.sira for x in katalog] + [0]) + 1, aktif=True)); db.commit()
                    log_action(db, "Oluşturma", detay=f"Proses adımı: {ad}")
                    flash(f"'{ad}' eklendi.", "success")
            elif a in ("yukari", "asagi", "pasif"):
                x = db.get(ProsesAdimTanim, f.get("adim_id", type=int) or 0)
                if x:
                    if a == "pasif":
                        x.aktif = not (x.aktif is not False)
                    else:
                        L = sorted(katalog, key=lambda z: (z.sira, z.id)); i = L.index(x); j = i - 1 if a == "yukari" else i + 1
                        if 0 <= j < len(L):
                            L[i].sira, L[j].sira = L[j].sira, L[i].sira
                            for n, z in enumerate(L, 1):
                                pass
                    db.commit()
            return redirect(url_for("urun_proses_kutuphane", adim=f.get("adim_id") or "", std=_std(f.get("std"))))
        aid = request.args.get("adim", type=int) or (katalog[0].id if katalog else 0)
        adim = db.get(ProsesAdimTanim, aid)
        std = _std(request.args.get("std"))
        fm = _std_sorgu(db.query(ProsesFmeaSatiri).filter_by(adim_tanim_id=aid, urun_proses_id=None), std).order_by(ProsesFmeaSatiri.sira, ProsesFmeaSatiri.id).all() if adim else []
        kp = db.query(ProsesKpSatiri).filter_by(adim_tanim_id=aid, urun_proses_id=None).order_by(ProsesKpSatiri.sira, ProsesKpSatiri.id).all() if adim else []
        K = FC.kriterler()
        return render_template("urun_proses_kutuphane.html", katalog=katalog, adim=adim, fmea=[_fmea_json(r) for r in fm], kp=[_kp_json(r) for r in kp],
                               yonetici=_yonetici(), puanlar=FC.puan_listesi(std), turler=TURLER, std=std, standartlar=FC.STANDARTLAR,
                               ipuclari=K["ipuclari"].get("P", []) if std == "GSI" else [], kriterler=K["kriterler"].get("P", {}) if std == "GSI" else {})
    finally:
        db.close()


# ─────────────────────────── ürünün proses akışı ───────────────────────────
def _urun_verisi(db, u):
    katalog = {k.id: k for k in _katalog(db)}
    std = _std(u)
    adimlar = db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).order_by(UrunProsesAdimi.sira, UrunProsesAdimi.id).all()
    out = []
    for a in adimlar:
        t = katalog.get(a.adim_tanim_id)
        genel = _std_sorgu(db.query(ProsesFmeaSatiri).filter_by(adim_tanim_id=a.adim_tanim_id, urun_proses_id=None), std).order_by(ProsesFmeaSatiri.sira, ProsesFmeaSatiri.id).all()
        ozel = _std_sorgu(db.query(ProsesFmeaSatiri).filter_by(adim_tanim_id=a.adim_tanim_id, urun_proses_id=u.id), std).order_by(ProsesFmeaSatiri.sira, ProsesFmeaSatiri.id).all()
        kp = db.query(ProsesKpSatiri).filter_by(adim_tanim_id=a.adim_tanim_id, urun_proses_id=u.id).order_by(ProsesKpSatiri.sira, ProsesKpSatiri.id).all()
        sablon = db.query(ProsesKpSatiri).filter_by(adim_tanim_id=a.adim_tanim_id, urun_proses_id=None).all()
        var = {r.kaynak_id for r in kp if r.kaynak_id}
        eksik = sum(1 for s in sablon if s.id not in var)
        satir_ap = [_fmea_json(r) for r in genel + ozel]
        out.append({"uid": a.id, "tanim_id": a.adim_tanim_id, "ad": t.ad if t else "?", "genel": [_fmea_json(r) for r in genel], "ozel": [_fmea_json(r) for r in ozel],
                    "kp": [_kp_json(r) for r in kp], "eksik_sablon": eksik, "ozet": FC.ozet([{k: x[k] for k in FMEA_METIN + FMEA_PUAN} for x in satir_ap], "P", std),
                    "olcum_hazir": sum(1 for r in kp if r.tur != "Görsel" and (r.alt_limit is not None or r.ust_limit is not None))})
    return out


def _adim_ekle(db, u, tanim_id):
    sira = (db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).count()) + 1
    a = UrunProsesAdimi(urun_proses_id=u.id, adim_tanim_id=tanim_id, sira=sira)
    db.add(a); db.flush()
    _sablondan_getir(db, u, tanim_id)
    return a


def _sablondan_getir(db, u, tanim_id):
    """Adımın kontrol planı şablon satırlarından ürüne henüz kopyalanmamış olanları kopyalar. (kopyalanan sayısı)"""
    var = {r.kaynak_id for r in db.query(ProsesKpSatiri).filter_by(adim_tanim_id=tanim_id, urun_proses_id=u.id).all() if r.kaynak_id}
    n = 0
    for s in db.query(ProsesKpSatiri).filter_by(adim_tanim_id=tanim_id, urun_proses_id=None).order_by(ProsesKpSatiri.sira, ProsesKpSatiri.id).all():
        if s.id in var:
            continue
        db.add(ProsesKpSatiri(adim_tanim_id=tanim_id, urun_proses_id=u.id, kaynak_id=s.id, sira=s.sira, karakteristik=s.karakteristik, tur=s.tur, birim=s.birim,
                              nominal=s.nominal, alt_limit=s.alt_limit, ust_limit=s.ust_limit, ozel=s.ozel, olcum_araci=s.olcum_araci, numune=s.numune,
                              siklik=s.siklik, yontem=s.yontem, reaksiyon=s.reaksiyon, fmea_satir_id=s.fmea_satir_id))
        n += 1
    return n


@app.route("/urun-proses/<int:uid>", methods=["GET", "POST"])
@login_required
def urun_proses_detay(uid):
    db = get_db()
    try:
        u = db.get(UrunProsesi, uid)
        if not u:
            abort(404)
        if request.method == "POST":
            if not _yonetici():
                abort(403)
            f = request.form
            a = f.get("action")
            git = url_for("urun_proses_detay", uid=u.id)
            if a == "adim_ekle":
                tid = f.get("adim_tanim_id", type=int)
                if tid and db.get(ProsesAdimTanim, tid):
                    ua = _adim_ekle(db, u, tid); db.commit()
                    log_action(db, "Güncelleme", detay=f"Ürün prosesi {u.urun_adi}: adım eklendi (otomatik FMEA + kontrol planı şablonu)")
                    return redirect(git + f"#adim-{ua.id}")
            elif a in ("yukari", "asagi", "adim_sil", "sablon"):
                ua = db.get(UrunProsesAdimi, f.get("uadim", type=int) or 0)
                if not ua or ua.urun_proses_id != u.id:
                    abort(404)
                if a == "adim_sil":
                    if db.query(OlcumKaydi).filter_by(adim_id=ua.id).count():
                        flash("Bu adımda ölçü kaydı var; silinemez.", "error")
                    else:
                        for r in db.query(ProsesKpSatiri).filter_by(adim_tanim_id=ua.adim_tanim_id, urun_proses_id=u.id).all():
                            db.delete(r)
                        for r in db.query(ProsesFmeaSatiri).filter_by(adim_tanim_id=ua.adim_tanim_id, urun_proses_id=u.id).all():
                            db.delete(r)
                        db.delete(ua); db.commit()
                elif a == "sablon":
                    n = _sablondan_getir(db, u, ua.adim_tanim_id); db.commit()
                    flash(f"Şablondan {n} satır getirildi." if n else "Eksik şablon satırı yok.", "success")
                else:
                    L = db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).order_by(UrunProsesAdimi.sira, UrunProsesAdimi.id).all()
                    i = L.index(ua); j = i - 1 if a == "yukari" else i + 1
                    if 0 <= j < len(L):
                        L[i], L[j] = L[j], L[i]
                    for n, x in enumerate(L, 1):
                        x.sira = n
                    db.commit()
            elif a == "bilgi":
                if (f.get("urun_adi") or "").strip():
                    u.urun_adi = f.get("urun_adi").strip()[:200]
                u.urun_kodu = (f.get("urun_kodu") or "").strip()[:80] or None
                u.aciklama = (f.get("aciklama") or "").strip()[:500] or None
                u.aktif = f.get("aktif") == "1"
                ns = f.get("fmea_standart")
                if ns in FC.STANDARTLAR and ns != _std(u):
                    if db.query(ProsesFmeaSatiri).filter(ProsesFmeaSatiri.urun_proses_id == u.id).count():
                        flash("Ürüne özgü FMEA satırı varken FMEA standardı değiştirilemez.", "error")
                    else:
                        u.fmea_standart = ns
                db.commit(); flash("Kaydedildi.", "success")
            return redirect(git)
        K = FC.kriterler()
        katalog = [k for k in _katalog(db) if k.aktif is not False]
        veri = _urun_verisi(db, u)
        return render_template("urun_proses_detay.html", u=u, adimlar=veri, katalog=katalog, yonetici=_yonetici(), turler=TURLER,
                               puanlar=FC.puan_listesi(_std(u)), ipuclari=K["ipuclari"].get("P", []) if _std(u) == "GSI" else [],
                               kriterler=K["kriterler"].get("P", {}) if _std(u) == "GSI" else {}, std=_std(u), standartlar=FC.STANDARTLAR)
    finally:
        db.close()


# ─────────────────────────── satır API'leri (FMEA / kontrol planı) ───────────────────────────
@app.route("/urun-proses/api/fmea", methods=["POST"])
@login_required
def urun_proses_api_fmea():
    if not _yonetici():
        abort(403)
    f = request.form
    db = get_db()
    try:
        a = f.get("action")
        if a == "ekle":
            tid = f.get("adim_tanim_id", type=int)
            if not tid or not db.get(ProsesAdimTanim, tid):
                abort(404)
            up = f.get("urun_proses_id", type=int) or None
            pu = db.get(UrunProsesi, up) if up else None
            std = _std(pu) if pu else _std(f.get("standart"))
            n = _std_sorgu(db.query(ProsesFmeaSatiri).filter_by(adim_tanim_id=tid, urun_proses_id=up), std).count()
            r = ProsesFmeaSatiri(adim_tanim_id=tid, urun_proses_id=up, standart=std, sira=n + 1)
            db.add(r); db.commit()
            return jsonify(ok=True, satir=_fmea_json(r))
        r = db.get(ProsesFmeaSatiri, f.get("rid", type=int) or 0)
        if not r:
            abort(404)
        if a == "sil":
            db.query(ProsesKpSatiri).filter_by(fmea_satir_id=r.id).update({"fmea_satir_id": None})
            db.delete(r); db.commit()
            return jsonify(ok=True, silindi=True)
        alan, deger = f.get("alan"), (f.get("deger") or "").strip()
        if alan in FMEA_PUAN:
            if not FC.dogrula_puan("P", alan[0], deger, _std(r.standart)):
                return jsonify(ok=False, hata="Bu puan seçilemez."), 400
            setattr(r, alan, int(deger) if deger else None)
        elif alan in FMEA_METIN:
            setattr(r, alan, deger[:4000] or None)
        else:
            abort(400)
        db.commit()
        return jsonify(ok=True, satir=_fmea_json(r))
    finally:
        db.close()


@app.route("/urun-proses/api/kp", methods=["POST"])
@login_required
def urun_proses_api_kp():
    if not _yonetici():
        abort(403)
    f = request.form
    db = get_db()
    try:
        a = f.get("action")
        if a == "ekle":
            tid = f.get("adim_tanim_id", type=int)
            if not tid or not db.get(ProsesAdimTanim, tid):
                abort(404)
            up = f.get("urun_proses_id", type=int) or None
            n = db.query(ProsesKpSatiri).filter_by(adim_tanim_id=tid, urun_proses_id=up).count()
            r = ProsesKpSatiri(adim_tanim_id=tid, urun_proses_id=up, sira=n + 1, tur="Ölçü", numune=1)
            db.add(r); db.commit()
            return jsonify(ok=True, satir=_kp_json(r))
        r = db.get(ProsesKpSatiri, f.get("rid", type=int) or 0)
        if not r:
            abort(404)
        if a == "sil":
            if db.query(OlcumKaydi).filter_by(kp_satir_id=r.id).count():
                return jsonify(ok=False, hata="Bu karakteristikte ölçü kaydı var; silinemez."), 400
            db.delete(r); db.commit()
            return jsonify(ok=True, silindi=True)
        alan, deger = f.get("alan"), (f.get("deger") or "").strip()
        try:
            if alan in KP_SAYI:
                setattr(r, alan, _sayi(deger))
            elif alan == "numune":
                v = int(deger) if deger else None
                if v is not None and not 1 <= v <= 50:
                    raise ValueError("Numune 1–50")
                r.numune = v
            elif alan == "fmea_satir_id":
                r.fmea_satir_id = int(deger) if deger else None
            elif alan == "tur":
                if deger and deger not in TURLER:
                    raise ValueError("Geçersiz tür")
                r.tur = deger or None
            elif alan == "ozel":
                if deger not in ("", "CC", "SC"):
                    raise ValueError("CC / SC")
                r.ozel = deger or None
            elif alan in KP_METIN:
                setattr(r, alan, deger[:KP_METIN[alan]] or None)
            else:
                abort(400)
        except ValueError as e:
            return jsonify(ok=False, hata=str(e)), 400
        if r.alt_limit is not None and r.ust_limit is not None and r.alt_limit > r.ust_limit:
            db.rollback()
            return jsonify(ok=False, hata="Alt limit üst limitten büyük olamaz."), 400
        db.commit()
        return jsonify(ok=True, satir=_kp_json(r))
    finally:
        db.close()


# ─────────────────────────── ölçü girişi ───────────────────────────
def _sonuc(v, alt, ust):
    if v is None:
        return None
    if alt is not None and v < alt:
        return "NOK"
    if ust is not None and v > ust:
        return "NOK"
    return "OK"


@app.route("/urun-proses/<int:uid>/olcum", methods=["GET", "POST"])
@login_required
def urun_proses_olcum(uid):
    db = get_db()
    try:
        u = db.get(UrunProsesi, uid)
        if not u:
            abort(404)
        adimlar = db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).order_by(UrunProsesAdimi.sira, UrunProsesAdimi.id).all()
        katalog = {k.id: k for k in _katalog(db)}
        ua = db.get(UrunProsesAdimi, request.values.get("adim", type=int) or 0)
        if ua and ua.urun_proses_id != u.id:
            ua = None
        if not ua and len(adimlar) == 1:
            ua = adimlar[0]
        kp = []
        if ua:
            kp = [r for r in db.query(ProsesKpSatiri).filter_by(adim_tanim_id=ua.adim_tanim_id, urun_proses_id=u.id).order_by(ProsesKpSatiri.sira, ProsesKpSatiri.id).all() if r.tur != "Görsel"]
        olculebilir = [r for r in kp if r.alt_limit is not None or r.ust_limit is not None]
        if request.method == "POST":
            if not _yazabilir():
                abort(403)
            if not ua:
                abort(400)
            f = request.form
            girisler, hata, nok = [], [], 0
            for r in olculebilir:
                for n in range(1, (r.numune or 1) + 1):
                    ham = (f.get(f"v_{r.id}_{n}") or "").strip()
                    if not ham:
                        continue
                    try:
                        v = _sayi(ham)
                    except ValueError:
                        hata.append(f"{r.karakteristik}: '{ham}' sayı değil")
                        continue
                    girisler.append((r, n, v))
            if hata or not girisler:
                flash("; ".join(hata) if hata else "En az bir ölçü değeri girin.", "error")
                return render_template("urun_proses_olcum.html", **_olcum_ctx(db, u, ua, adimlar, katalog, olculebilir, kp, f))
            grup = f"OL-{datetime.now():%y%m%d%H%M%S}-{session.get('user_id') or 0}"
            for r, n, v in girisler:
                s = _sonuc(v, r.alt_limit, r.ust_limit)
                nok += s == "NOK"
                db.add(OlcumKaydi(grup=grup, urun_proses_id=u.id, adim_id=ua.id, kp_satir_id=r.id, vardiya=(f.get("vardiya") or "")[:10] or None,
                                  is_emri=(f.get("is_emri") or "").strip()[:40] or None, lot_no=(f.get("lot_no") or "").strip()[:40] or None,
                                  numune_no=n, deger=v, sonuc=s, olcen_id=session.get("user_id")))
            db.commit()
            log_action(db, "Oluşturma", detay=f"Ölçü girişi {grup}: {u.urun_adi} / {katalog[ua.adim_tanim_id].ad} — {len(girisler)} değer, {nok} NOK")
            flash(f"{len(girisler)} ölçü kaydedildi." + (f" {nok} değer tolerans dışı (NOK) — Uygun Olmayan Ürün / DÖF süreci başlatın." if nok else " Tümü uygun (OK)."),
                  "warning" if nok else "success")
            return redirect(url_for("urun_proses_olcum", uid=u.id, adim=ua.id, son=grup))
        return render_template("urun_proses_olcum.html", **_olcum_ctx(db, u, ua, adimlar, katalog, olculebilir, kp, {}, son=request.args.get("son")))
    finally:
        db.close()


def _olcum_ctx(db, u, ua, adimlar, katalog, olculebilir, kp, form, son=None):
    sonlar = []
    if son:
        q = {r.id: r for r in olculebilir}
        for k in db.query(OlcumKaydi).filter_by(grup=son).order_by(OlcumKaydi.kp_satir_id, OlcumKaydi.numune_no).all():
            sonlar.append({"kar": q[k.kp_satir_id].karakteristik if k.kp_satir_id in q else k.kp_satir_id, "no": k.numune_no, "deger": k.deger, "sonuc": k.sonuc})
    tanimsiz = [r for r in kp if r.alt_limit is None and r.ust_limit is None]
    return dict(u=u, ua=ua, adimlar=[{"id": a.id, "ad": katalog[a.adim_tanim_id].ad} for a in adimlar], satirlar=olculebilir, tanimsiz=tanimsiz,
                form=form, sonlar=sonlar, yazabilir=_yazabilir(), vardiyalar=("1", "2", "3"), son=son)


# ─────────────────────────── geçmiş / istatistik / Excel ───────────────────────────
def _istatistik(vals, alt, ust):
    n = len(vals)
    if not n:
        return {}
    ort = statistics.fmean(vals)
    s = statistics.stdev(vals) if n > 1 else 0.0
    d = {"n": n, "ort": ort, "min": min(vals), "max": max(vals), "s": s, "cp": None, "cpk": None}
    if n >= 5 and s > 0:
        if alt is not None and ust is not None:
            d["cp"] = (ust - alt) / (6 * s)
        eller = []
        if ust is not None:
            eller.append((ust - ort) / (3 * s))
        if alt is not None:
            eller.append((ort - alt) / (3 * s))
        d["cpk"] = min(eller) if eller else None
    return d


@app.route("/urun-proses/<int:uid>/olcumler")
@login_required
def urun_proses_olcumler(uid):
    db = get_db()
    try:
        u = db.get(UrunProsesi, uid)
        if not u:
            abort(404)
        katalog = {k.id: k for k in _katalog(db)}
        adimlar = {a.id: a for a in db.query(UrunProsesAdimi).filter_by(urun_proses_id=u.id).all()}
        kps = {r.id: r for r in db.query(ProsesKpSatiri).filter_by(urun_proses_id=u.id).all()}
        adim_f = request.args.get("adim", type=int)
        kayit = db.query(OlcumKaydi).filter_by(urun_proses_id=u.id).order_by(OlcumKaydi.tarih.desc(), OlcumKaydi.id.desc()).all()
        if adim_f:
            kayit = [k for k in kayit if k.adim_id == adim_f]
        if request.args.get("excel"):
            return _excel(u, kayit, kps, adimlar, katalog, db)
        stat = []
        for kid, r in kps.items():
            v = [k.deger for k in kayit if k.kp_satir_id == kid and k.deger is not None]
            if v:
                ua = next((a for a in adimlar.values() if a.adim_tanim_id == r.adim_tanim_id), None)
                s = _istatistik(v, r.alt_limit, r.ust_limit)
                s.update({"kar": r.karakteristik, "birim": r.birim or "", "alt": r.alt_limit, "ust": r.ust_limit, "adim": katalog[r.adim_tanim_id].ad if r.adim_tanim_id in katalog else "",
                          "nok": sum(1 for k in kayit if k.kp_satir_id == kid and k.sonuc == "NOK")})
                stat.append(s)
        kullanici = {x.id: x.ad_soyad for x in db.query(User).all()}
        return render_template("urun_proses_olcumler.html", u=u, kayit=[(k, kps.get(k.kp_satir_id), katalog.get(adimlar[k.adim_id].adim_tanim_id).ad if k.adim_id in adimlar else "",
                                                                         kullanici.get(k.olcen_id, "—")) for k in kayit[:400]], toplam=len(kayit), stat=stat,
                               adimlar=[(a.id, katalog[a.adim_tanim_id].ad) for a in sorted(adimlar.values(), key=lambda z: z.sira)], adim_f=adim_f)
    finally:
        db.close()


def _excel(u, kayit, kps, adimlar, katalog, db):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    kullanici = {x.id: x.ad_soyad for x in db.query(User).all()}
    wb = Workbook(); ws = wb.active; ws.title = "Ölçüler"
    ws.append([f"{u.musteri} — {u.urun_adi} {u.urun_kodu or ''} — ölçü kayıtları ({date.today():%d.%m.%Y})"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Tarih", "Adım", "Karakteristik", "Birim", "Nominal", "Alt limit", "Üst limit", "Numune", "Değer", "Sonuç", "Vardiya", "İş emri", "Lot", "Ölçen", "Kontrol grubu"])
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
    for k in kayit:
        r = kps.get(k.kp_satir_id); a = adimlar.get(k.adim_id)
        ws.append([k.tarih, katalog[a.adim_tanim_id].ad if a else "", r.karakteristik if r else "", r.birim if r else "", r.nominal if r else None, r.alt_limit if r else None,
                   r.ust_limit if r else None, k.numune_no, k.deger, k.sonuc, k.vardiya, k.is_emri, k.lot_no, kullanici.get(k.olcen_id, ""), k.grup])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY HH:MM"
        if k.sonuc == "NOK":
            ws.cell(ws.max_row, 10).fill = PatternFill("solid", fgColor="F4B6B6")
    for i, w in enumerate([17, 20, 30, 7, 9, 9, 9, 8, 10, 8, 8, 12, 12, 18, 22], 1):
        ws.column_dimensions[ws.cell(2, i).column_letter].width = w
    ws.freeze_panes = "A3"; ws.auto_filter.ref = f"A2:O{ws.max_row}"
    bio = io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"Ölçüler {u.urun_adi} {date.today():%Y%m%d}.xlsx"[:150],
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/urun-proses/<int:uid>/kontrol-plani")
@login_required
def urun_proses_kontrol_plani(uid):
    """Ürünün kontrol planı + FMEA özeti — Excel (Control Plan düzeni)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    db = get_db()
    try:
        u = db.get(UrunProsesi, uid)
        if not u:
            abort(404)
        wb = Workbook(); ws = wb.active; ws.title = "Kontrol Planı"
        ws.append([f"KONTROL PLANI — {u.musteri} — {u.urun_adi} {u.urun_kodu or ''} — {date.today():%d.%m.%Y}"]); ws["A1"].font = Font(bold=True, size=12)
        ws.append(["Proses adımı", "Karakteristik", "Tür", "Özel", "Birim", "Nominal", "Alt limit", "Üst limit", "Ölçüm aracı", "Numune", "Sıklık", "Kontrol yöntemi", "Reaksiyon planı", "İlgili FMEA (hata modu / AP)"])
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(wrap_text=True)
        for a in _urun_verisi(db, u):
            fm = {x["id"]: x for x in a["genel"] + a["ozel"]}
            for r in a["kp"]:
                f = fm.get(r["fmea_satir_id"]) if r["fmea_satir_id"] != "" else None
                ws.append([a["ad"], r["karakteristik"], r["tur"], r["ozel"], r["birim"], r["nominal"], r["alt_limit"], r["ust_limit"], r["olcum_araci"], r["numune"], r["siklik"], r["yontem"],
                           r["reaksiyon"], f"{f['mode']} / {f['ap_ad']}" if f else ""])
        for i, w in enumerate([20, 30, 8, 6, 7, 9, 9, 9, 22, 8, 18, 28, 30, 30], 1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        ws2 = wb.create_sheet("FMEA (P-FMEA)")
        ws2.append(["Proses adımı", "Kapsam", "Fonksiyon", "Hata etkisi", "S", "Hata modu", "Hata nedeni", "Önleyici aksiyon", "O", "Tespit aksiyonu", "D", "AP", "RPN (bilgi)", "Optimizasyon aksiyonu", "O′", "D′", "AP′"])
        for c in ws2[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E78")
        renk = {"H": "F4B6B6", "M": "FFE699", "L": "C6E0B4"}
        for a in _urun_verisi(db, u):
            for kap, L in (("Genel", a["genel"]), ("Ürüne özgü", a["ozel"])):
                for x in L:
                    ws2.append([a["ad"], kap, x["fn"], x["eff"], x["S"], x["mode"], x["cause"], x["prev"], x["O"], x["det"], x["D"], x["ap_ad"], x["rpn"], x["act"], x["O2"], x["D2"], x["ap2_ad"]])
                    for col, ap in ((12, x["ap"]), (17, x["ap2"])):
                        if ap:
                            ws2.cell(ws2.max_row, col).fill = PatternFill("solid", fgColor=renk[ap])
        bio = io.BytesIO(); wb.save(bio); bio.seek(0)
        log_action(db, "Görüntüleme", detay=f"Kontrol planı / FMEA Excel: {u.urun_adi}")
        return send_file(bio, as_attachment=True, download_name=f"Kontrol Planı {u.urun_adi} {date.today():%Y%m%d}.xlsx"[:150],
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        db.close()
